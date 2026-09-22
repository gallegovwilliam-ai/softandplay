from django.shortcuts import render, HttpResponse
from django.contrib.auth.decorators import login_required
from django.http import Http404
from django.views.generic import TemplateView
# Create your views here.
from .models import CabezeraImpre ,Impresion, Carton, CabezeraImpre75 ,Impresion75, Carton75
from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from clientes.models import Cliente
from partidas.models import Partida, Partida75
from pagos.reconciliation import register_sale, register_sale75
from .services import allocate_cartons
import datetime
import os
from openpyxl import load_workbook
from django.utils import timezone

def xls(request, nee=None):
    if request.method != "POST" or not request.user.is_superuser:
        raise Http404
    n = (int(nee)*1000)-1000
    file = os.environ.get("SOFTANDPLAY_CARTONES_XLSX")
    if not file:
        return HttpResponse("Falta SOFTANDPLAY_CARTONES_XLSX", status=503)
    #file = "/opt/multibingos/app/DB.xlsx"
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.active
    if n < 0 or n + 1000 > sheet.max_row:
        return HttpResponse('Bloque de 1000 cartones fuera del rango del archivo', status=400)
    with transaction.atomic():
        for i in range(1000):
            c = Carton.objects.create(
            	l1_1 = str(int(sheet.cell(row=i+n+1, column=1).value)),
            	l1_2 = str(int(sheet.cell(row=i+n+1, column=2).value)), 
            	l1_3 = str(int(sheet.cell(row=i+n+1, column=3).value)), 
            	l1_4 = str(int(sheet.cell(row=i+n+1, column=4).value)), 
            	l2_1 = str(int(sheet.cell(row=i+n+1, column=5).value)),
            	l2_2 = str(int(sheet.cell(row=i+n+1, column=6).value)), 
            	l2_3 = str(int(sheet.cell(row=i+n+1, column=7).value)), 
            	l2_4 = str(int(sheet.cell(row=i+n+1, column=8).value)), 
            	l3_1 = str(int(sheet.cell(row=i+n+1, column=9).value)),
            	l3_2 = str(int(sheet.cell(row=i+n+1, column=10).value)), 
            	l3_3 = str(int(sheet.cell(row=i+n+1, column=11).value)), 
            	l3_4 = str(int(sheet.cell(row=i+n+1, column=12).value)), 
            	l4_1 = str(int(sheet.cell(row=i+n+1, column=13).value)),
            	l4_2 = str(int(sheet.cell(row=i+n+1, column=14).value)), 
            	l4_3 = str(int(sheet.cell(row=i+n+1, column=15).value)), 
            	l4_4 = str(int(sheet.cell(row=i+n+1, column=16).value)), 
            	)
    return HttpResponse('loto mexicana')


#bingo 75

def xls75(request, nee=None):
    if request.method != "POST" or not request.user.is_superuser:
        raise Http404
    n = (int(nee)*1000)-1000
    file = os.environ.get("SOFTANDPLAY_CARTONES75_XLSX")
    if not file:
        return HttpResponse("Falta SOFTANDPLAY_CARTONES75_XLSX", status=503)
    #file = "/opt/multibingos/app/DB75.xlsx"
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.active
    if n < 0 or n + 1000 > sheet.max_row:
        return HttpResponse('Bloque de 1000 cartones fuera del rango del archivo', status=400)
    with transaction.atomic():
        for i in range(1000):
            c = Carton75.objects.create(
            	l1_1 = str(int(sheet.cell(row=i+n+1, column=1).value)),
            	l1_2 = str(int(sheet.cell(row=i+n+1, column=2).value)), 
            	l1_3 = str(int(sheet.cell(row=i+n+1, column=3).value)), 
            	l1_4 = str(int(sheet.cell(row=i+n+1, column=4).value)), 
            	l1_5 = str(int(sheet.cell(row=i+n+1, column=5).value)), 
            	l2_1 = str(int(sheet.cell(row=i+n+1, column=6).value)),
            	l2_2 = str(int(sheet.cell(row=i+n+1, column=7).value)), 
            	l2_3 = str(int(sheet.cell(row=i+n+1, column=8).value)), 
            	l2_4 = str(int(sheet.cell(row=i+n+1, column=9).value)), 
            	l2_5 = str(int(sheet.cell(row=i+n+1, column=10).value)),
            	l3_1 = str(int(sheet.cell(row=i+n+1, column=11).value)),
            	l3_2 = str(int(sheet.cell(row=i+n+1, column=12).value)), 
            	l3_3 = str(int(sheet.cell(row=i+n+1, column=13).value)), 
            	l3_4 = str(int(sheet.cell(row=i+n+1, column=14).value)), 
            	l3_5 = str(int(sheet.cell(row=i+n+1, column=15).value)), 
            	l4_1 = str(int(sheet.cell(row=i+n+1, column=16).value)),
            	l4_2 = str(int(sheet.cell(row=i+n+1, column=17).value)), 
            	l4_3 = str(int(sheet.cell(row=i+n+1, column=18).value)), 
            	l4_4 = str(int(sheet.cell(row=i+n+1, column=19).value)), 
            	l4_5 = str(int(sheet.cell(row=i+n+1, column=20).value)), 
            	l5_1 = str(int(sheet.cell(row=i+n+1, column=21).value)),
            	l5_2 = str(int(sheet.cell(row=i+n+1, column=22).value)), 
            	l5_3 = str(int(sheet.cell(row=i+n+1, column=23).value)), 
            	l5_4 = str(int(sheet.cell(row=i+n+1, column=24).value)), 
            	l5_5 = str(int(sheet.cell(row=i+n+1, column=25).value)), 
            	)
    return HttpResponse('bingo 75')


#loto mexicana
class Imprimir(TemplateView):
    template_name='impresiones.html'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['partidas'] =  Partida.objects.filter(inicio=False,termino=False,fecha__gte=timezone.localdate()).order_by('fecha')
        return context
Inprimir = Imprimir.as_view()

#bingo 75
class Imprimir75(TemplateView):
    template_name='impresiones75.html'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['partidas'] =  Partida75.objects.filter(inicio=False,termino=False,fecha__gte=timezone.localdate()).order_by('fecha')
        return context
Inprimir75 = Imprimir75.as_view()

#loto mexicana
@login_required
def tikect(request):
    if request.method != 'POST':
        return HttpResponse(status=405)
    if not (request.user.is_superuser or request.user.userprofile.user_type == 'Revendedor'):
        raise Http404
    try:
        cantidad = int(request.POST.get('v', '0'))
    except (TypeError, ValueError):
        return HttpResponse('<h2>Cantidad invalida</h2>', status=400)
    if cantidad < 1 or cantidad > 100:
        return HttpResponse('<h2>La cantidad debe estar entre 1 y 100</h2>', status=400)

    with transaction.atomic():
        partida = Partida.objects.select_for_update().filter(pk=request.POST.get('partida')).first()
        if not partida:
            return HttpResponse('<h2>Aun No hay Partidas creadas para el dia</h2>', status=404)
        if partida.inicio or partida.termino:
            return HttpResponse('<h2>Partida no disponible</h2>', status=409)

        idcliente = request.POST.get('c', '')
        cliente = None
        if idcliente:
            cliente = Cliente.objects.filter(pk=idcliente, user=request.user).first()
        if cliente is None:
            cliente = Cliente.objects.create(nombre=request.POST.get('n','')[:200], telefono=request.POST.get('t','')[:50], user=request.user)

        cabezera = CabezeraImpre.objects.create(fecha=timezone.now().date(), cliente=cliente, partida=partida, user=request.user, cantidad=cantidad, verficado=True)
        obj = allocate_cartons(partida, cabezera, cantidad)
        register_sale(cabezera, request.user, request=request)

    if request.POST['v'] == '1':
        return render(request,'imprimir.html',{'obj':obj})
    elif request.POST['v'] == '3':
        return render(request,'imprimir3.html',{'obj':obj})
    else:
        return render(request,'imprimir2.html',{'obj':obj})        

#bingo 75
@login_required
def tikect75(request):
    if request.method != 'POST':
        return HttpResponse(status=405)
    if not (request.user.is_superuser or request.user.userprofile.user_type == 'Revendedor'):
        raise Http404
    try:
        cantidad = int(request.POST.get('v', '0'))
    except (TypeError, ValueError):
        return HttpResponse('<h2>Cantidad invalida</h2>', status=400)
    if cantidad < 1 or cantidad > 100:
        return HttpResponse('<h2>La cantidad debe estar entre 1 y 100</h2>', status=400)

    with transaction.atomic():
        partida = Partida75.objects.select_for_update().filter(pk=request.POST.get('partida')).first()
        if not partida:
            return HttpResponse('<h2>Aun No hay Partidas creadas para el dia</h2>', status=404)
        if partida.inicio or partida.termino:
            return HttpResponse('<h2>Partida no disponible</h2>', status=409)

        idcliente = request.POST.get('c', '')
        cliente = None
        if idcliente:
            cliente = Cliente.objects.filter(pk=idcliente, user=request.user).first()
        if cliente is None:
            cliente = Cliente.objects.create(nombre=request.POST.get('n','')[:200], telefono=request.POST.get('t','')[:50], user=request.user)

        cabezera = CabezeraImpre75.objects.create(fecha=timezone.now().date(), cliente=cliente, partida=partida, user=request.user, cantidad=cantidad, verficado=True)
        obj = allocate_cartons(partida, cabezera, cantidad, is75=True)
        register_sale75(cabezera, request.user, request=request)

    if request.POST['v'] == '1':
        return render(request,'imprimir75.html',{'obj':obj})
    elif request.POST['v'] == '3':
        return render(request,'imprimir375.html',{'obj':obj})
    else:
        return render(request,'imprimir275.html',{'obj':obj})        
